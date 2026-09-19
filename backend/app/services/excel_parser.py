from openpyxl import load_workbook
from io import BytesIO
from rapidfuzz import process, fuzz
import re, unicodedata

def clean(v): 
    return re.sub(r"\s+"," ",str(v or "").replace("\u00a0"," ")).strip()

def norm(v):
    s=unicodedata.normalize('NFKD',clean(v)).encode('ascii','ignore').decode().lower()

    return re.sub(r'\s+',' ',re.sub(r'[^a-z0-9 ]','',s)).strip()

def parse_workbook(raw:bytes):
    wb=load_workbook(BytesIO(raw),data_only=True)
    sheets=[]; people={}; positions=set(); marks=set()
    stats={'valid_rows':0,'empty_rows':0,'warnings':[]}

    for ws in wb.worksheets:
        songs=[]; current=None; start_col=None

        for c in range(2,ws.max_column+1):
            h=clean(ws.cell(1,c).value); pos=clean(ws.cell(2,c).value)

            if h:
                current=h; start_col=c; songs.append({'name':h,'columns':[]})

            if current and pos:
                songs[-1]['columns'].append({'column':c,'position':pos})
                positions.add(pos)

        sheet_songs=[]

        for song in songs:
            assignments=[]
            cols=song['columns']

            for r in range(3,ws.max_row+1):
                person=clean(ws.cell(r,1).value)

                if not person:
                    stats['empty_rows']+=1
                    continue

                row_marked=False

                for col in cols:
                    val=clean(ws.cell(r,col['column']).value)

                    if val:
                        row_marked=True
                        marks.add(val)
                        assignments.append({'person':person,'position':col['position'],'mark':val})
                        people.setdefault(norm(person),person)

                if row_marked:
                    stats['valid_rows']+=1
                else:
                    stats['warnings'].append(f"Hoja '{ws.title}', canción '{song['name']}', fila {r}: persona '{person}' sin ninguna marca de participación.")

            if assignments:
                sheet_songs.append({'name':song['name'],'assignments':assignments})

        sheets.append({'name':ws.title,'songs':sheet_songs})

    return {'sheets':sheets,'people':sorted(people.values()),'positions':sorted(positions),'marks':sorted(marks),'stats':stats}

def match_people(source_names, canonical_names):
    out=[]

    choices={norm(x):x for x in canonical_names}
    keys=list(choices)

    for src in source_names:
        key=norm(src)

        if key in choices: 
            continue

        best=process.extractOne(key,keys,scorer=fuzz.ratio) if keys else None

        if best and best[1]>=70: 
            out.append({'source':src,'candidate':choices[best[0]],'score':round(best[1],1)})

    return out


def detect_duplicates(names):
    grouped={}
    for x in names:
        grouped.setdefault(norm(x),[]).append(x)
    exact=[sorted(v) for v in grouped.values() if len(v)>1]
    fuzzy=[]
    seen=set()
    items=sorted({norm(x):x for x in names}.values())

    for i,a in enumerate(items):
        for b in items[i+1:]:
            if norm(a)==norm(b):
                continue

            score=fuzz.ratio(norm(a),norm(b))

            if score>=88 and (norm(a),norm(b)) not in seen:
                seen.add((norm(a),norm(b)))
                fuzzy.append({'names':[a,b],'score':round(score,1)})

    return {'exact':exact,'similar':fuzzy}