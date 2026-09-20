from fastapi.testclient import TestClient
from io import BytesIO
from openpyxl import Workbook
from app.main import app

client=TestClient(app)

def test_health():
    assert client.get('/api/health').json()=={'status':'ok'}

def test_default_templates_are_seeded():
    r=client.get('/api/marimba-templates')

    assert r.status_code==200

    names={t['name'] for t in r.json()}

    assert {'Marimba tenor','Marimba grande','Marimbito'}<=names

def test_composition_roundtrip_keeps_configurable_marimba():
    p=client.post('/api/projects',json={'name':'Proyecto test'}).json()
    data={'elements':[{
        'id':'m1','type':'marimba','name':'Marimba Grande 1','x':10,'y':20,
        'width':380,'height':150,'rotation':0,'scaleX':1,'scaleY':1,
        'positions':[{'id':'p1','type':'Primera','personId':None},{'id':'p2','type':'Centro','personId':None}]
    }]}
    c=client.post('/api/compositions',json={'project_id':p['id'],'name':'Comp 1','data':data})

    assert c.status_code==200

    cid=c.json()['id']
    got=client.get(f'/api/compositions/{cid}').json()

    assert got['data']==data

    upd=client.put(f'/api/compositions/{cid}',json={'project_id':p['id'],'name':'Comp 1b','data':data})

    assert upd.status_code==200 and upd.json()['name']=='Comp 1b'

def test_import_confirm_creates_project_with_songs_and_assignments():
    payload={'project_name':'Concierto test','sheets':[{'name':'Hoja1','songs':[
        {'name':'Canción 1','assignments':[{'person':'Tiffany Salazar','position':'Centro','mark':'X'}]}
    ]}]}
    r=client.post('/api/imports/confirm',json=payload)

    assert r.status_code==200

    proj=client.get(f"/api/projects/{r.json()['id']}").json()
    a=proj['songs'][0]['assignments'][0]

    assert a['person']=='Tiffany Salazar' and a['position']=='Centro' and a['mark']=='X'

def _make_xlsx(rows):
    wb=Workbook()
    ws=wb.active
    ws['A1']='Nombre'; ws['B1']='Canción 1'
    ws['B2']='Primera'

    for i,(name,mark) in enumerate(rows,start=3):
        ws.cell(row=i,column=1,value=name)
        ws.cell(row=i,column=2,value=mark)

    buf=BytesIO()
    wb.save(buf)

    return buf.getvalue()

def test_preview_valid_excel():
    content=_make_xlsx([('Tiffany Salazar','X'),('María López','X')])
    r=client.post('/api/imports/preview',files={'file':('personas.xlsx',content)})

    assert r.status_code==200

    d=r.json()

    assert set(d['people'])>=({'Tiffany Salazar','María López'})
    assert d['positions']==['Primera']
    assert d['stats']['valid_rows']==2
    assert 'duplicates' in d

def test_preview_rejects_invalid_extension():
    r=client.post('/api/imports/preview',files={'file':('datos.csv',b'a,b')})

    assert r.status_code==400
    assert 'extensión' in r.json()['detail'].lower()

def test_preview_rejects_empty_file():
    r=client.post('/api/imports/preview',files={'file':('vacio.xlsx',b'')})

    assert r.status_code==400
    assert 'vacío' in r.json()['detail'].lower()

def test_preview_rejects_corrupt_file():
    r=client.post('/api/imports/preview',files={'file':('roto.xlsx',b'contenido que no es un zip')})

    assert r.status_code==422

def test_preview_rejects_oversized_file():
    big=b'\x00'*(10*1024*1024+1)
    r=client.post('/api/imports/preview',files={'file':('grande.xlsx',big)})

    assert r.status_code==413
    assert 'tamaño' in r.json()['detail'].lower()

def test_preview_excel_without_valid_rows():
    wb=Workbook()
    ws=wb.active
    ws['A1']='ID'; ws['B1']='Stock Actual'; ws['A2']=1
    buf=BytesIO(); wb.save(buf)
    r=client.post('/api/imports/preview',files={'file':('sin_datos.xlsx',buf.getvalue())})

    assert r.status_code==422
    assert 'no contiene filas válidas' in r.json()['detail'].lower()

def test_confirm_into_existing_project_avoids_duplicate_people():
    p=client.post('/api/projects',json={'name':'Proyecto import'}).json()
    sheets=[{'name':'H','songs':[{'name':'C','assignments':[
        {'person':'Ana García','position':'Tenor','mark':'X'},
        {'person':'ana  garcía','position':'Bajo','mark':'X'},
    ]}]}]
    r1=client.post('/api/imports/confirm',json={'project_id':p['id'],'source_filename':'a.xlsx','sheets':sheets})

    assert r1.status_code==200 and r1.json()['id']==p['id']

    r2=client.post('/api/imports/confirm',json={'project_id':p['id'],'source_filename':'b.xlsx','sheets':sheets})

    assert r2.status_code==200

    people=client.get('/api/people').json()
    anas=[x for x in people if 'ana' in x['name'].lower()]

    assert len(anas)==1

    proj=client.get(f"/api/projects/{p['id']}").json()
    positions=sorted({a['position'] for s in proj['songs'] for a in s['assignments']})

    assert positions==['Bajo','Tenor']
    assert proj['project']['source_filename']=='a.xlsx'

def test_confirm_unknown_project_returns_404():
    r=client.post('/api/imports/confirm',json={'project_id':999999,'sheets':[]})

    assert r.status_code==404

def test_song_crud_and_integrity():
    p=client.post('/api/projects',json={'name':'Concierto 20 de febrero 2026'}).json()
    s1=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Luna de Xelajú'})

    assert s1.status_code==200 and s1.json()['name']=='Luna de Xelajú'

    s2=client.post(f"/api/projects/{p['id']}/songs",json={'name':'El Grito'}).json()
    s3=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Noche de Luna'}).json()
    songs=client.get(f"/api/projects/{p['id']}/songs").json()
    
    assert [s['name'] for s in songs]==['Luna de Xelajú','El Grito','Noche de Luna']

    dup=client.post(f"/api/projects/{p['id']}/songs",json={'name':'luna de xelajú'})

    assert dup.status_code==409

    empty=client.post(f"/api/projects/{p['id']}/songs",json={'name':'   '})

    assert empty.status_code==400

    ren=client.patch(f"/api/songs/{s2['id']}",json={'name':'El Grito (final)'})

    assert ren.status_code==200 and ren.json()['name']=='El Grito (final)'

    other=client.post('/api/projects',json={'name':'Otro proyecto'}).json()
    oc=client.post(f"/api/projects/{other['id']}/songs",json={'name':'Canción ajena'}).json()
    cross=client.post('/api/compositions',json={'project_id':p['id'],'song_id':oc['id'],'name':'X','data':{}})

    assert cross.status_code==400

    ghost=client.post('/api/compositions',json={'project_id':p['id'],'song_id':999999,'name':'X','data':{}})

    assert ghost.status_code==404

    nop=client.post('/api/compositions',json={'project_id':999999,'song_id':None,'name':'X','data':{}})

    assert nop.status_code==404

    d=client.delete(f"/api/songs/{s3['id']}")

    assert d.status_code==200
    assert client.get(f"/api/songs/{s3['id']}").status_code==404

    songs=client.get(f"/api/projects/{p['id']}/songs").json()

    assert [s['name'] for s in songs]==['Luna de Xelajú','El Grito (final)']

def test_duplicate_composition_independent_copy():
    p=client.post('/api/projects',json={'name':'Proyecto dup'}).json()
    s1=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Luna de Xelajú'}).json()
    s2=client.post(f"/api/projects/{p['id']}/songs",json={'name':'El Grito'}).json()
    data={'elements':[{'id':'m1','type':'marimba','name':'Marimba A','x':1,'y':2,'width':380,'height':150,
        'rotation':15,'scaleX':1.2,'scaleY':1.2,
        'positions':[{'id':'p1','type':'Primera','personId':7},{'id':'p2','type':'Segunda','personId':None}]}]}
    c=client.post('/api/compositions',json={'project_id':p['id'],'song_id':s1['id'],'name':'Luna de Xelajú','data':data}).json()
    d=client.post(f"/api/compositions/{c['id']}/duplicate",json={'song_id':s2['id'],'name':'El Grito - distribución'})

    assert d.status_code==200

    copy=d.json()

    assert copy['id']!=c['id'] and copy['song_id']==s2['id']
    assert copy['data']==data

    mod=dict(data)
    mod['elements']=[dict(data['elements'][0])]
    mod['elements'][0]['name']='MODIFICADA'

    client.put(f"/api/compositions/{copy['id']}",json={'project_id':p['id'],'song_id':s2['id'],'name':'El Grito - distribución','data':mod})

    orig=client.get(f"/api/compositions/{c['id']}").json()
    got=client.get(f"/api/compositions/{copy['id']}").json()

    assert orig['data']['elements'][0]['name']=='Marimba A'
    assert got['data']['elements'][0]['name']=='MODIFICADA'

    dup_default=client.post(f"/api/compositions/{c['id']}/duplicate",json={})
    
    assert dup_default.status_code==200 and dup_default.json()['name'].endswith('(copia)')

    bad=client.post(f"/api/compositions/{c['id']}/duplicate",json={'song_id':999999})

    assert bad.status_code==404

def test_duplicate_rejects_song_from_other_project():
    p1=client.post('/api/projects',json={'name':'P1'}).json()
    p2=client.post('/api/projects',json={'name':'P2'}).json()
    s2=client.post(f"/api/projects/{p2['id']}/songs",json={'name':'Canción P2'}).json()
    c=client.post('/api/compositions',json={'project_id':p1['id'],'name':'C P1','data':{}}).json()
    r=client.post(f"/api/compositions/{c['id']}/duplicate",json={'song_id':s2['id']})

    assert r.status_code==400

def test_get_song_returns_404_for_missing():
    assert client.get('/api/songs/999999').status_code==404

def test_rename_composition():
    p=client.post('/api/projects',json={'name':'Proyecto rename'}).json()
    s=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Luna'}).json()
    c=client.post('/api/compositions',json={'project_id':p['id'],'song_id':s['id'],'name':'Luna - Ensayo','data':{}}).json()

    ren=client.patch(f"/api/compositions/{c['id']}",json={'name':'Luna - Final'})
    assert ren.status_code==200
    assert ren.json()['name']=='Luna - Final'

    got=client.get(f"/api/compositions/{c['id']}").json()
    assert got['name']=='Luna - Final'

    bad=client.patch(f"/api/compositions/{c['id']}",json={'name':'   '})
    assert bad.status_code==400

def test_duplicate_composition_endpoint():
    p=client.post('/api/projects',json={'name':'Proyecto dup test'}).json()
    s=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Canción 1'}).json()
    c=client.post('/api/compositions',json={'project_id':p['id'],'song_id':s['id'],'name':'Original','data':{'elements':[]}}).json()

    dup=client.post(f"/api/compositions/{c['id']}/duplicate",json={'name':'Duplicada'})
    assert dup.status_code==200
    res=dup.json()
    assert res['id']!=c['id']
    assert res['name']=='Duplicada'
    assert res['song_id']==s['id']

def test_delete_composition():
    p=client.post('/api/projects',json={'name':'Proyecto delete'}).json()
    c=client.post('/api/compositions',json={'project_id':p['id'],'name':'Para borrar','data':{}}).json()

    d=client.delete(f"/api/compositions/{c['id']}")
    assert d.status_code==200
    assert d.json()['deleted'] is True
    assert d.json()['id']==c['id']

    assert client.get(f"/api/compositions/{c['id']}").status_code==404

def test_composition_independent_after_duplicate():
    p=client.post('/api/projects',json={'name':'Proyecto indep'}).json()
    s=client.post(f"/api/projects/{p['id']}/songs",json={'name':'Canción'}).json()
    data={'elements':[{'id':'m1','type':'marimba','name':'Marimba 1','positions':[]}]}
    c1=client.post('/api/compositions',json={'project_id':p['id'],'song_id':s['id'],'name':'C1','data':data}).json()

    c2=client.post(f"/api/compositions/{c1['id']}/duplicate",json={'name':'C2'}).json()
    assert c2['id']!=c1['id']

    # Update c2, verify c1 untouched
    mod2={'elements':[{'id':'m1','type':'marimba','name':'Marimba MOD C2','positions':[]}]}
    client.put(f"/api/compositions/{c2['id']}",json={'project_id':p['id'],'song_id':s['id'],'name':'C2','data':mod2})
    assert client.get(f"/api/compositions/{c1['id']}").json()['data']['elements'][0]['name']=='Marimba 1'
    assert client.get(f"/api/compositions/{c2['id']}").json()['data']['elements'][0]['name']=='Marimba MOD C2'

    # Update c1, verify c2 untouched
    mod1={'elements':[{'id':'m1','type':'marimba','name':'Marimba MOD C1','positions':[]}]}
    client.put(f"/api/compositions/{c1['id']}",json={'project_id':p['id'],'song_id':s['id'],'name':'C1','data':mod1})
    assert client.get(f"/api/compositions/{c1['id']}").json()['data']['elements'][0]['name']=='Marimba MOD C1'
    assert client.get(f"/api/compositions/{c2['id']}").json()['data']['elements'][0]['name']=='Marimba MOD C2'

def test_composition_belongs_to_correct_song():
    p1=client.post('/api/projects',json={'name':'P1 Song Check'}).json()
    p2=client.post('/api/projects',json={'name':'P2 Song Check'}).json()
    s1=client.post(f"/api/projects/{p1['id']}/songs",json={'name':'Song 1'}).json()
    s2=client.post(f"/api/projects/{p2['id']}/songs",json={'name':'Song 2'}).json()

    # Attempt cross-project song assignment on create -> 400
    r1=client.post('/api/compositions',json={'project_id':p1['id'],'song_id':s2['id'],'name':'Comp Bad','data':{}})
    assert r1.status_code==400

    # Valid creation with song in same project
    r2=client.post('/api/compositions',json={'project_id':p1['id'],'song_id':s1['id'],'name':'Comp Good','data':{}})
    assert r2.status_code==200
    cid=r2.json()['id']

    # Attempt cross-project song update on patch -> 400
    r3=client.patch(f"/api/compositions/{cid}",json={'song_id':s2['id']})
    assert r3.status_code==400

    # Song list includes this composition
    song_info=client.get(f"/api/songs/{s1['id']}").json()
    comp_ids=[c['id'] for c in song_info['compositions']]
    assert cid in comp_ids

def test_composition_not_found_handling():
    assert client.get('/api/compositions/999999').status_code==404
    assert client.patch('/api/compositions/999999',json={'name':'X'}).status_code==404
    assert client.delete('/api/compositions/999999').status_code==404
    assert client.post('/api/compositions/999999/duplicate',json={'name':'X'}).status_code==404
    assert client.put('/api/compositions/999999',json={'project_id':1,'name':'X'}).status_code==404

def test_delete_composition_does_not_delete_song_or_people():
    payload={'project_name':'Concierto No Delete Cascade','sheets':[{'name':'Hoja1','songs':[
        {'name':'Luna de Xelajú','assignments':[{'person':'Carlos Méndez','position':'Primera','mark':'X'}]}
    ]}]}
    proj=client.post('/api/imports/confirm',json=payload).json()
    pid=proj['id']
    song_data=client.get(f"/api/projects/{pid}").json()['songs'][0]
    sid=song_data['id']

    # Create composition for this song
    comp=client.post('/api/compositions',json={'project_id':pid,'song_id':sid,'name':'Comp Test Delete','data':{}}).json()
    cid=comp['id']

    # Delete composition
    del_res=client.delete(f"/api/compositions/{cid}")
    assert del_res.status_code==200

    # Check composition is gone
    assert client.get(f"/api/compositions/{cid}").status_code==404

    # Song still exists
    song_chk=client.get(f"/api/songs/{sid}")
    assert song_chk.status_code==200
    assert song_chk.json()['name']=='Luna de Xelajú'
    assert len(song_chk.json()['assignments'])==1
    assert song_chk.json()['assignments'][0]['person']=='Carlos Méndez'

    # People still exists
    people=client.get('/api/people').json()
    assert any(p['name']=='Carlos Méndez' for p in people)

    # Project still exists
    assert client.get(f"/api/projects/{pid}").status_code==200

def test_compatibility_with_old_composition():
    # Legacy data without 'locked', without scaleX/scaleY or with minimal fields
    p=client.post('/api/projects',json={'name':'Legacy Proj'}).json()
    legacy_data={
        'elements': [
            {'id':'m_old','type':'marimba','name':'Marimba Antigua','x':50,'y':50,'width':380,'height':150,'positions':['Primera','Bajo']},
            {'id':'p_old','type':'person','name':'Persona Antigua','personId':1,'x':60,'y':60}
        ]
    }
    c=client.post('/api/compositions',json={'project_id':p['id'],'name':'Legacy Comp','data':legacy_data})
    assert c.status_code==200
    cid=c.json()['id']

    got=client.get(f"/api/compositions/{cid}").json()
    assert got['data']==legacy_data
